"""Exportação integral, sob demanda, da conferência de baixas SETTA."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import json

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TABLES = (
    ("RESUMO_OP", "ops"),
    ("PRODUTOS", "items"),
    ("LOTES", "lots"),
    ("RELATORIO_MES", "general"),
    ("SALDOS_ENDERECOS", "stock"),
    ("MOVIMENTOS_KARDEX", "movements"),
)


def _safe_cell(value):
    if value is None:
        return None
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str) and value and value[0] in ("=", "+", "-", "@"):
        # Conteúdos de planilhas externas não podem tornar-se fórmulas Excel.
        return "'" + value
    return value


def _frame(values) -> pd.DataFrame:
    if isinstance(values, pd.DataFrame):
        result = values.copy()
    else:
        result = pd.DataFrame(values or [])
    for col in result.columns:
        result[col] = result[col].map(_safe_cell)
    return result


def export_complete_report(model: dict) -> bytes:
    """Retorna Excel com TODOS os registros, independente de filtro/paginação."""
    output = BytesIO()
    meta = model.get("source_meta") or {}
    sources = [
        {"fonte": key, **value}
        for key, value in meta.items()
    ]
    sources += [
        {"fonte": "KARDEX", **entry}
        for entry in model.get("history_meta") or []
    ]
    coverage = model.get("coverage") or {}
    summary = pd.DataFrame([
        {"indicador": "Referência da conferência", "valor": model.get("as_of")},
        {"indicador": "Armazéns selecionados", "valor": ", ".join(model.get("warehouses") or ())},
        {"indicador": "KARDEX disponível", "valor": "SIM" if model.get("history_available") else "NÃO"},
        {"indicador": "Período inicial", "valor": coverage.get("inicio")},
        {"indicador": "Período final", "valor": coverage.get("fim")},
        {"indicador": "Movimentos após a base", "valor": model.get("movimentos_apos_base", 0)},
        {"indicador": "Linhas com saldo negativo", "valor": model.get("negative_stock_rows", 0)},
        {"indicador": "Versão da regra", "valor": model.get("version", "")},
        {"indicador": "Itens MES", "valor": len(model.get("general") or [])},
        {"indicador": "Movimentos vinculados à OP", "valor": len(model.get("movements") or [])},
    ])
    with pd.ExcelWriter(output, engine="openpyxl") as excel:
        summary.to_excel(excel, sheet_name="PARAMETROS", index=False)
        for sheet, key in TABLES:
            frame = _frame(model.get(key))
            if frame.empty and len(frame.columns) == 0:
                frame = pd.DataFrame(columns=["SEM REGISTROS"])
            frame.to_excel(excel, sheet_name=sheet, index=False)
        _frame(sources).to_excel(excel, sheet_name="FONTES", index=False)
        _frame([{"intervalo_inicio": start, "intervalo_fim": end}
                for start, end in coverage.get("intervalos") or []]).to_excel(
                    excel, sheet_name="COBERTURA", index=False)
        for worksheet in excel.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = PatternFill("solid", fgColor="111827")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            worksheet.row_dimensions[1].height = 29
            for col_cells in worksheet.columns:
                letter = get_column_letter(col_cells[0].column)
                max_width = max(
                    [len(str(cell.value or "")) for cell in list(col_cells)[:200]] or [10]
                )
                worksheet.column_dimensions[letter].width = min(max(max_width + 3, 13), 55)
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        cell.number_format = "#,##0.######"
    return output.getvalue()
