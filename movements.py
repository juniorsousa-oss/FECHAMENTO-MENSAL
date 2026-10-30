"""KARDEX: histórico físico, importações idempotentes e agrupamento neutro por TES."""
from __future__ import annotations

import base64
import gzip
import hashlib
import io
import json
import re
import unicodedata
import xml.etree.ElementTree as ET
from functools import lru_cache
from collections import Counter, defaultdict, deque
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import pandas as pd
import kardex_history as history

RULES_FILE = Path(__file__).parent / "config" / "movement_rules.json"
FORMAT = "SETTA_KARDEX_IMPORT_V1"
HISTORY_FORMAT = "SETTA_KARDEX_HISTORY_V1"
ADJUSTMENT = "AJUSTE DE ESTOQUE"
FLOW = "FLUXO DE ESTOQUE"
PRODUCTION = "PRODUÇÃO INTERNA"
CLOSING = "AJUSTE DE FECHAMENTO DE CHAPAS E BARRAS"
CLOSING_RULE = "020_520_FECHAMENTO_CHAPAS_BARRAS"
NUMERIC = {
    "custo_medio_inicio_arquivo", "saldo_inicio_arquivo", "valor_saldo_inicio_arquivo",
    "qtd_entrada", "valor_entrada", "custo_medio_movimento", "qtd_saida", "valor_saida",
    "saldo_pos_lancamento", "valor_saldo_pos_lancamento",
}
COLUMNS = {
    "A": "codigo", "B": "descricao", "C": "um", "D": "tipo", "E": "grupo",
    "F": "custo_medio_inicio_arquivo", "G": "saldo_inicio_arquivo", "H": "valor_saldo_inicio_arquivo",
    "I": "posicao_ipi", "J": "endereco_relatorio", "L": "armazem", "M": "tes", "N": "cf",
    "O": "documento", "Q": "qtd_entrada", "S": "valor_entrada", "U": "custo_medio_movimento",
    "W": "qtd_saida", "Y": "valor_saida", "AA": "saldo_pos_lancamento", "AC": "valor_saldo_pos_lancamento",
    "AE": "ae_original",
}
RAW_KEYS = set(COLUMNS.values()) | {
    "data", "arquivo_origem", "linha_origem", "arquivo_sha256",
    "periodo_arquivo_inicio", "periodo_arquivo_fim",
}
IDENTITY_KEYS = ("data", "codigo", "armazem", "um", "tes", "cf", "documento", "qtd_entrada",
                 "valor_entrada", "qtd_saida", "valor_saida", "saldo_pos_lancamento",
                 "valor_saldo_pos_lancamento", "ae_original")
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


TES_CATALOG_FILE = Path(__file__).parent / "config" / "tes_cadastro_protheus_20261022.b64"
TES_CATALOG_VERSION = "TES-PROTHEUS-20261022+TES-DESCRICOES-PRIORITARIAS-20261022"
TES_OVERRIDES_FILE = Path(__file__).parent / "config" / "tes_descricoes_prioritarias.json"


@lru_cache(maxsize=1)
def tes_reference() -> dict[str, str]:
    """342 descrições do campo Txt Padrao do cadastro TES Protheus.

    A representação gzip/base64 mantém pequeno o catálogo originalmente extraído
    da planilha TES(1).xltx. Não modifica o banco nem lançamentos históricos.
    """
    lines = [
        line.strip() for line in TES_CATALOG_FILE.read_text(encoding="utf-8").splitlines()
        if line.strip() and not line.startswith("#")
    ]
    catalog = json.loads(gzip.decompress(base64.b64decode("".join(lines), validate=True)))
    if len(catalog) != 342:
        raise ValueError("Catálogo TES inconsistente: deveriam existir 342 códigos.")
    return catalog


@lru_cache(maxsize=1)
def tes_overrides() -> dict[str, str]:
    """Nove descrições confirmadas que prevalecem sobre o catálogo-base."""
    payload = json.loads(TES_OVERRIDES_FILE.read_text(encoding="utf-8"))
    tes = payload["tes"]
    expected = {"010", "020", "030", "040", "520", "540", "550", "499", "999"}
    if set(tes) != expected or any(not isinstance(value, str) or not value.strip() for value in tes.values()):
        raise ValueError("Descrições TES prioritárias inconsistentes.")
    return tes


def rules() -> dict:
    """Uma descrição por TES; exceções do usuário substituem o cadastro-base."""
    config = json.loads(RULES_FILE.read_text(encoding="utf-8"))
    descriptions = tes_reference().copy()
    descriptions.update(tes_overrides())
    return {
        "versao": config["versao"] + "+" + TES_CATALOG_VERSION,
        "modo": "CADASTRO_TES_SOMENTE",
        "tes": {code: {"descricao": value} for code, value in descriptions.items()},
    }


def normalized(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").upper())
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).split())


def decimal(value) -> Decimal:
    if value in (None, ""):
        return Decimal(0)
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Número inválido: {value!r}") from exc
    if not result.is_finite():
        raise ValueError("Número não finito no relatório.")
    return result


def identifier(value, width=0) -> str:
    text = "" if value is None else str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]
    return text.zfill(width) if text.isdigit() else text


def _date(value, date1904=False) -> str:
    if value in (None, ""):
        raise ValueError("Data da movimentação não informada.")
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    text = str(value or "").strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}", text):
        return date.fromisoformat(text[:10]).isoformat()
    for fmt in ("%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    try:
        origin = datetime(1904, 1, 1) if date1904 else datetime(1899, 12, 30)
        return (origin + timedelta(days=float(decimal(value)))).date().isoformat()
    except Exception as exc:
        raise ValueError(f"Data da movimentação inválida: {value!r}") from exc


def clean_record(source: dict) -> dict:
    row = {key: source.get(key) for key in RAW_KEYS}
    if row.get("ae_original") is None:
        row["ae_original"] = source.get("vinculo", "")
    row["ae_original"] = str(row["ae_original"] or "")  # original spacing is meaningful evidence
    row["codigo"] = identifier(row["codigo"], 8)
    row["tes"] = identifier(row["tes"], 3)
    row["armazem"] = identifier(row["armazem"])
    row["data"] = _date(row["data"])
    row["linha_origem"] = int(row.get("linha_origem") or 0)
    for key in NUMERIC:
        row[key] = None if row.get(key) in (None, "") else str(decimal(row[key]))
    for key in ("descricao", "um", "tipo", "grupo", "cf", "documento", "endereco_relatorio"):
        row[key] = str(row.get(key) or "")
    if not row["codigo"] or not row["tes"] or not row["armazem"]:
        raise ValueError("Movimentação sem produto, TES ou armazém.")
    return row


def fingerprint(row: dict) -> str:
    parts = [str(decimal(row.get(key)).normalize()) if key in NUMERIC else str(row.get(key) or "")
             for key in IDENTITY_KEYS]
    return hashlib.sha256(json.dumps(parts, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()


def _xml_rows(raw: bytes):
    import posixpath
    with ZipFile(io.BytesIO(raw)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 700_000_000:
            raise ValueError("Arquivo descompactado muito grande. Divida o relatório por período.")
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        links = {r.get("Id"): r.get("Target") for r in rels}
        props = workbook.find("s:workbookPr", NS)
        epoch1904 = props is not None and props.get("date1904") in ("1", "true")
        strings = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            strings = ["".join(t.text or "" for t in si.findall(".//s:t", NS)) for si in shared]
        for sheet in workbook.find("s:sheets", NS):
            rid = sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            target = links.get(rid)
            if not target:
                continue
            target = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            with archive.open(target) as stream:
                for _, node in ET.iterparse(stream, events=("end",)):
                    if node.tag.rsplit("}", 1)[-1] != "row":
                        continue
                    cells = {}
                    for c in node:
                        address = c.get("r", "")
                        column = re.sub(r"\d", "", address)
                        kind = c.get("t")
                        value = c.findtext("s:v", namespaces=NS)
                        if kind == "inlineStr":
                            value = "".join(t.text or "" for t in c.findall(".//s:t", NS))
                        elif kind == "s" and value is not None:
                            value = strings[int(value)]
                        if value not in (None, ""):
                            cells[column] = value
                    if cells:
                        yield sheet.get("name", ""), int(node.get("r", "0")), cells, epoch1904
                    node.clear()


def parse_report(raw: bytes, file_name: str) -> dict:
    """Reads the ERP's physical AE column; opening-only rows are not movements."""
    digest = hashlib.sha256(raw).hexdigest()
    if file_name.lower().endswith(".zip"):
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 700_000_000:
                raise ValueError("Base histórica muito grande para esta importação.")
            manifest = json.loads(archive.read("manifesto.json"))
            if manifest.get("formato") != HISTORY_FORMAT:
                raise ValueError("O ZIP não corresponde a uma base histórica KARDEX.")
            rows = [clean_record(json.loads(line)) for line in archive.read("movimentacoes.jsonl").splitlines() if line.strip()]
        if not rows and not manifest.get("coberturas"):
            raise ValueError("Base histórica sem movimentações.")
        meta = {"arquivo_nome": file_name, "sha256": digest, "lancamentos": len(rows),
                "periodo_inicio": manifest.get("periodo_inicio") or min(r["data"] for r in rows),
                "periodo_fim": manifest.get("periodo_fim") or max(r["data"] for r in rows),
                "coberturas_arquivo": manifest.get("coberturas"),
                "linhas_sem_movimentacao": 0}
        return {"meta": meta, "rows": rows}
    rows, errors, skipped = [], [], 0
    sheets_with_header = set()
    for sheet, line, cells, epoch1904 in _xml_rows(raw):
        if normalized(cells.get("A")) in ("CODIGO", "COD. PRODUTO") and normalized(cells.get("M")) == "TES":
            if "OP" not in normalized(cells.get("AE")) or "DATA" not in normalized(cells.get("K")):
                raise ValueError("Layout KARDEX diferente: confira Data em K, TES em M e vínculo original em AE.")
            sheets_with_header.add(sheet)
            continue
        if sheet not in sheets_with_header:
            continue
        if not cells.get("K") and not cells.get("M"):
            skipped += 1
            continue
        try:
            row = {name: cells.get(col) for col, name in COLUMNS.items()}
            row.update({"data": _date(cells.get("K"), epoch1904), "arquivo_origem": file_name,
                        "arquivo_sha256": digest, "linha_origem": line})
            rows.append(clean_record(row))
        except (ValueError, TypeError) as exc:
            errors.append(f"{sheet}, linha {line}: {exc}")
    if not sheets_with_header:
        raise ValueError("Cabeçalho KARDEX não encontrado. Importe o relatório original de movimentação dos produtos.")
    if errors:
        raise ValueError(f"{len(errors)} linhas incompletas. O arquivo não foi importado. " + "; ".join(errors[:5]))
    if not rows:
        raise ValueError("O relatório não possui movimentações válidas.")
    start, end = min(r["data"] for r in rows), max(r["data"] for r in rows)
    # The ERP title has no period; a valid period in its export filename can include non-working days.
    dates = re.findall(r"(\d{2})[./](\d{2})(?:[./](\d{4}))?", file_name)
    declared_period = False
    if len(dates) == 2:
        try:
            first = date(int(dates[0][2] or start[:4]), int(dates[0][1]), int(dates[0][0])).isoformat()
            last = date(int(dates[1][2] or end[:4]), int(dates[1][1]), int(dates[1][0])).isoformat()
            if first <= start <= end <= last:
                start, end = first, last
                declared_period = True
        except ValueError:
            pass
    for row in rows:
        row["periodo_arquivo_inicio"], row["periodo_arquivo_fim"] = start, end
    return {"meta": {"arquivo_nome": file_name, "sha256": digest, "lancamentos": len(rows),
                     "periodo_inicio": start, "periodo_fim": end, "linhas_sem_movimentacao": skipped,
                     "periodo_informado": declared_period}, "rows": rows}


def pack_import(parsed: dict) -> dict:
    rows = [clean_record(row) for row in parsed["rows"]]
    encoded = json.dumps(rows, ensure_ascii=False, separators=(",", ":")).encode()
    return {"format": FORMAT, "meta": dict(parsed["meta"]), "encoding": "gzip+base64",
            "payload": base64.b64encode(gzip.compress(encoded, compresslevel=6, mtime=0)).decode("ascii")}


def unpack_import(pack: dict) -> dict:
    if pack.get("format") != FORMAT or pack.get("encoding") != "gzip+base64":
        raise ValueError("Formato da importação histórica inválido.")
    raw = gzip.decompress(base64.b64decode(pack["payload"], validate=True))
    rows = [clean_record(row) for row in json.loads(raw)]
    if len(rows) != int(pack["meta"]["lancamentos"]):
        raise ValueError("Contagem de registros diferente da importação salva.")
    return {"meta": pack["meta"], "rows": rows}


def merge_imports(imports: list[dict]) -> tuple[list[dict], int]:
    snapshot = history_snapshot(imports)
    return snapshot["rows"], snapshot["duplicates"]


def history_snapshot(imports: list[dict]) -> dict:
    snapshot = history.resolve(imports)
    rows, duplicates = _merge_active(snapshot["imports"])
    return {**snapshot, "rows": rows, "duplicates": duplicates}


def _merge_active(imports: list[dict]) -> tuple[list[dict], int]:
    """Keep max multiplicity across reports, never collapse repeated rows inside one report."""
    merged, seen, duplicates = [], set(), 0
    for imported in imports:
        occurrences = Counter()
        for raw in imported["rows"]:
            row = clean_record(raw)
            signature = fingerprint(row)
            occurrences[signature] += 1
            identity = f"{signature}:{occurrences[signature]}"
            if identity in seen:
                duplicates += 1
                continue
            seen.add(identity)
            row["id_registro"] = hashlib.sha256(identity.encode()).hexdigest()[:24]
            merged.append(row)
    return merged, duplicates


def classify(row: dict, config: dict) -> dict:
    """Enriquece o registro sem inferir grupo, condição, consumo ou tipo.

    A TES e a descrição vêm diretamente do cadastro Protheus. Não converter
    descrição em uma classificação de negócio nem usar regras antigas.
    """
    result = dict(row)
    code = row["tes"]
    tes_data = config["tes"].get(code, {})
    description = tes_data.get("descricao") or "NI - TES não localizada no cadastro"
    ae = str(row.get("ae_original") or "").strip()
    op = re.fullmatch(r"OP\s*(\d+)", ae, re.I)
    cc = re.fullmatch(r"CC\s*(\d+)", ae, re.I)
    result.update({
        "descricao_tes": description,
        "grupo_analise": "", "subgrupo_analise": "",
        "grupo_provisorio": False, "regra_classificacao": "",
        "motivo_classificacao": "", "material_h001": False,
        "op_identificada": "OP" + op.group(1) if op else "",
        "cc_identificado": cc.group(1) if cc else "",
        "tipo_vinculo_ae": "OP" if op else "CC" if ae.upper().startswith("CC") else "OUTRO",
        "enderecamento_compensado": False, "par_enderecamento": "",
        "enderecamento_candidato": False, "versao_regra": config["versao"],
    })
    return result

def _direction(row):
    incoming = decimal(row.get("qtd_entrada")) != 0 or decimal(row.get("valor_entrada")) != 0
    outgoing = decimal(row.get("qtd_saida")) != 0 or decimal(row.get("valor_saida")) != 0
    return "entrada" if incoming and not outgoing else "saida" if outgoing and not incoming else ""


def match_addresses(records: list[dict]) -> list[dict]:
    waiting = defaultdict(deque)
    candidates, pairs = defaultdict(lambda: {"entrada": [], "saida": []}), []
    for row in records:
        if row["tes"] not in ("499", "999"):
            continue
        direction = _direction(row)
        if not direction:
            continue
        common = tuple(row[k] for k in ("codigo", "armazem", "um", "documento", "data"))
        qty = decimal(row.get("qtd_" + direction)); amount = decimal(row.get("valor_" + direction))
        opposite = "saida" if direction == "entrada" else "entrada"
        other_tes = "999" if row["tes"] == "499" else "499"
        other_key = common + (qty, amount, other_tes, opposite)
        if waiting[other_key]:
            other = waiting[other_key].popleft()
            pid = "END-" + hashlib.sha256((row["id_registro"] + other["id_registro"]).encode()).hexdigest()[:12]
            for r in (row, other):
                r["enderecamento_compensado"] = True; r["par_enderecamento"] = pid
                r["motivo_enderecamento"] = "Par 499/999 com produto, armazém, documento, data, quantidade e valor coincidentes; OP não integra a chave."
        else:
            waiting[common + (qty, amount, row["tes"], direction)].append(row)
    for queue in waiting.values():
        for row in queue:
            if row["enderecamento_compensado"]:
                continue
            direction = _direction(row)
            key = tuple(row[k] for k in ("codigo", "armazem", "um", "documento", "data")) + (decimal(row.get("valor_" + direction)),)
            candidates[key][direction].append(row)
    for sides in candidates.values():
        if len(sides["entrada"]) != 1 or len(sides["saida"]) != 1:
            continue
        a, b = sides["entrada"][0], sides["saida"][0]
        if a["tes"] == b["tes"]:
            continue
        a["enderecamento_candidato"] = b["enderecamento_candidato"] = True
        pairs.append({"codigo": a["codigo"], "data": a["data"], "documento": a["documento"], "armazem": a["armazem"],
                      "um": a["um"], "quantidade_entrada": str(decimal(a["qtd_entrada"])),
                      "quantidade_saida": str(decimal(b["qtd_saida"])),
                      "diferenca": str(decimal(a["qtd_entrada"]) - decimal(b["qtd_saida"])),
                      "entrada_id": a["id_registro"], "saida_id": b["id_registro"]})
    return pairs


def build_model(imports: list[dict], config: dict | None = None) -> dict:
    cfg = dict(config or rules())
    snapshot = history_snapshot(imports)
    originals, duplicates = snapshot["rows"], snapshot["duplicates"]
    records = [classify(r, cfg) for r in originals]
    # Nenhum pareamento ou exclusão automática por suposta transferência:
    # nesta fase, cada lançamento do KARDEX permanece nos totais originais.
    address_pairs = []
    frame = pd.DataFrame(records)
    if not frame.empty:
        frame["competencia"] = frame["data"].str[:7]
        for key in NUMERIC:
            frame[key] = frame[key].map(lambda v: float(decimal(v)))
        frame["liquido_valor"] = frame["valor_entrada"] - frame["valor_saida"]
        frame["_search"] = (frame["codigo"] + " " + frame["descricao"] + " " + frame["descricao_tes"] + " " + frame["tes"] + " " + frame["documento"] + " " + frame["ae_original"]).map(normalized)
    token = hashlib.sha256((cfg["versao"] + history.VERSION + json.dumps(snapshot["coverage"], sort_keys=True)).encode()).hexdigest()
    return {"frame": frame, "rows": records, "raw_rows": originals, "duplicates": duplicates,
            "pares_enderecamento": address_pairs, "imports": [i["meta"] for i in imports], "versao": cfg["versao"], "token": token,
            "coverage_metas": snapshot["coverage"], "import_registry": snapshot["registry"]}


def effective(frame: pd.DataFrame) -> pd.DataFrame:
    """Etapa neutra: todos os lançamentos compõem os totais, sem ajustes automáticos."""
    return frame


def summarize(frame: pd.DataFrame, by: list[str]) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=by + ["lancamentos", "produtos", "tes_distintas", "valor_entrada", "valor_saida", "liquido_valor"])
    grouped = frame.groupby(by, dropna=False, sort=False).agg(
        lancamentos=("id_registro", "size"), produtos=("codigo", "nunique"), tes_distintas=("tes", "nunique"),
        valor_entrada=("valor_entrada", "sum"), valor_saida=("valor_saida", "sum"), liquido_valor=("liquido_valor", "sum"),
    ).reset_index()
    return grouped


def summarize_ops(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame()
    linked = frame.loc[frame["op_identificada"] != ""].copy()
    if linked.empty:
        return pd.DataFrame()
    linked["entrada_efetiva"] = linked["valor_entrada"]
    linked["saida_efetiva"] = linked["valor_saida"]
    return linked.groupby("op_identificada").agg(
        lancamentos=("id_registro", "size"), produtos=("codigo", "nunique"), documentos=("documento", "nunique"),
        primeira_data=("data", "min"), ultima_data=("data", "max"), enderecamentos=("enderecamento_compensado", "sum"),
        valor_entrada=("entrada_efetiva", "sum"), valor_saida=("saida_efetiva", "sum"),
    ).reset_index().sort_values("valor_saida", ascending=False, kind="stable").reset_index(drop=True)


EXPORT_HEADERS = ["Código", "Produto", "UM", "Descrição TES", "TES", "CF original", "Armazém", "Data", "Documento",
                  "Quantidade entrada", "Valor entrada", "Quantidade saída", "Valor saída", "Tipo vínculo AE", "Arquivo origem",
                  "Linha origem", "Endereço", "Grupo do arquivo", "Posição IPI", "Tipo do material", "CC", "Saldo após movimento",
                  "", "", "", "", "", "", "", "", "AE original"]


def ledger_export(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=EXPORT_HEADERS)
    data = {
        "Código": frame["codigo"], "Produto": frame["descricao"], "UM": frame["um"], "Descrição TES": frame["descricao_tes"],
        "TES": frame["tes"], "CF original": frame["cf"], "Armazém": frame["armazem"], "Data": pd.to_datetime(frame["data"]).dt.date,
        "Documento": frame["documento"], "Quantidade entrada": frame["qtd_entrada"], "Valor entrada": frame["valor_entrada"],
        "Quantidade saída": frame["qtd_saida"], "Valor saída": frame["valor_saida"], "Tipo vínculo AE": frame["tipo_vinculo_ae"],
        "Arquivo origem": frame["arquivo_origem"], "Linha origem": frame["linha_origem"], "Endereço": frame["endereco_relatorio"],
        "Grupo do arquivo": frame["grupo"], "Posição IPI": frame["posicao_ipi"], "Tipo do material": frame["tipo"],
        "CC": frame["cc_identificado"], "Saldo após movimento": frame["saldo_pos_lancamento"],
    }
    result = pd.DataFrame(data)
    for i in range(8):
        result[f"_reserva_{i}"] = ""
    result["AE original"] = frame["ae_original"]
    result.columns = EXPORT_HEADERS
    return result.reset_index(drop=True)


def export_excel(frame: pd.DataFrame, group_config: dict | None = None, *, summary_only: bool = False) -> bytes:
    """Exportação SETTA: Grupos > TES, OPs > baixas e registros como na tela.

    As linhas detalhadas são recolhidas pelo agrupamento nativo do Excel.
    A geração é streaming para suportar o histórico completo do KARDEX.
    Não altera registros nem presume TES como baixa: utiliza saída real.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    import kardex_group_config as gc
    import kardex_conditions as conditions

    config = gc.validated(group_config if group_config is not None else gc.empty())
    owner = gc.group_owner(config)
    active = conditions.classify(effective(frame), config)
    wb = Workbook(write_only=True)
    title_color = "111827"
    group_color = "E9EFF8"
    detail_color = "F8FAFC"
    total_color = "D9E6F5"
    money_format = '#,##0.00;[Red](#,##0.00)'
    quantity_format = '#,##0.000###;[Red](#,##0.000###)'
    counters = {}

    def start_sheet(name: str, headers: list[str], widths: list[int], *, outline: bool = False):
        sheet = wb.create_sheet(name)
        counters[name] = 0
        sheet.freeze_panes = "B2"
        sheet.sheet_view.showGridLines = False
        if outline:
            sheet.sheet_properties.outlinePr.summaryBelow = False
            sheet.sheet_properties.outlinePr.showSummaryRight = False
            sheet.sheet_view.showOutlineSymbols = True
            sheet.sheet_format.outlineLevelRow = 2
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        write(sheet, headers, style="header")
        return sheet

    def write(sheet, values, *, style="detail", level=0, collapsed=False,
              date_cols=(), money_cols=(), quantity_cols=()):
        counters[sheet.title] += 1
        number = counters[sheet.title]
        if level:
            sheet.row_dimensions[number].outlineLevel = level
            sheet.row_dimensions[number].hidden = True
        if collapsed:
            sheet.row_dimensions[number].collapsed = True
        if style == "header":
            sheet.row_dimensions[number].height = 30
        elif style in ("group", "total"):
            sheet.row_dimensions[number].height = 24
        filled = []
        for index, item in enumerate(values, 1):
            if item is not None and not isinstance(item, (list, dict)) and pd.isna(item):
                item = None
            if index in date_cols and isinstance(item, str) and item:
                item = date.fromisoformat(item[:10])
            cell = WriteOnlyCell(sheet, item)
            if isinstance(item, str):
                # Never interpret descriptions, product codes or documents as formulas.
                cell.data_type = "s"
            cell.font = Font(
                name="Arial", size=10,
                bold=style in ("header", "group", "total"),
                color="FFFFFF" if style == "header" else title_color,
            )
            if style == "header":
                cell.fill = PatternFill("solid", fgColor=title_color)
            elif style in ("group", "total"):
                cell.fill = PatternFill("solid", fgColor=group_color if style == "group" else total_color)
            elif style == "child":
                cell.fill = PatternFill("solid", fgColor=detail_color)
            cell.alignment = Alignment(
                vertical="center", wrap_text=style == "header",
                indent=1 if level and index == 1 else 0,
            )
            if index in date_cols and isinstance(item, (date, datetime)):
                cell.number_format = "dd/mm/yyyy"
            elif index in money_cols and isinstance(item, (int, float, Decimal)):
                cell.number_format = money_format
            elif index in quantity_cols and isinstance(item, (int, float, Decimal)):
                cell.number_format = quantity_format
            filled.append(cell)
        sheet.append(filled)

    # 01 — Grupo -> Subgrupo -> TES (ambos níveis recolhíveis no Excel).
    group_headers = ["Grupo / Subgrupo / TES", "Descrição TES", "TES distintas", "Movimentos",
                     "Produtos", "Entrada (R$)", "Saída (R$)", "Variação (R$)"]
    group_sheet = start_sheet(
        "Grupos e TES", group_headers, [43, 67, 18, 16, 15, 23, 23, 23], outline=True
    )
    if not active.empty:
        main_stats = summarize(active, ["grupo_usuario"])
        sub_stats = summarize(active, ["grupo_usuario", "subgrupo_usuario"])
        tes_stats = summarize(active, ["grupo_usuario", "subgrupo_usuario", "tes", "descricao_tes"])
        main_lookup = {r["grupo_usuario"]: r for r in main_stats.to_dict("records")}
        sub_lookup = {name: part.sort_values("subgrupo_usuario", kind="stable")
                      for name, part in sub_stats.groupby("grupo_usuario", sort=False)}
        detail_lookup = {(group, subgroup): part.sort_values(
            ["tes", "descricao_tes"], kind="stable"
        ) for (group, subgroup), part in
            tes_stats.groupby(["grupo_usuario", "subgrupo_usuario"], sort=False)}
        order = [g["nome"] for g in config["grupos"]]
        if gc.UNASSIGNED in main_lookup:
            order.append(gc.UNASSIGNED)
        for group_name in order:
            stat = main_lookup.get(group_name)
            if stat is None:
                continue
            write(group_sheet,
                  [group_name, "", int(stat["tes_distintas"]), int(stat["lancamentos"]),
                   int(stat["produtos"]), float(stat["valor_entrada"]),
                   float(stat["valor_saida"]), float(stat["liquido_valor"])],
                  style="group", collapsed=True, money_cols=(6, 7, 8))
            for subgroup in sub_lookup[group_name].itertuples(index=False):
                child = detail_lookup[(group_name, subgroup.subgrupo_usuario)]
                write(group_sheet, [
                    "SUBGRUPO · " + subgroup.subgrupo_usuario, "",
                    int(subgroup.tes_distintas), int(subgroup.lancamentos), int(subgroup.produtos),
                    float(subgroup.valor_entrada), float(subgroup.valor_saida),
                    float(subgroup.liquido_valor)
                ], style="group", level=1, collapsed=True, money_cols=(6, 7, 8))
                for entry in child.itertuples(index=False):
                    write(group_sheet, [
                        "TES " + str(entry.tes), str(entry.descricao_tes), 1,
                        int(entry.lancamentos), int(entry.produtos),
                        float(entry.valor_entrada), float(entry.valor_saida),
                        float(entry.liquido_valor)
                    ], style="child", level=2, money_cols=(6, 7, 8))
        write(group_sheet,
              ["TOTAL GERAL", "", int(active["tes"].nunique()), len(active),
               int(active["codigo"].nunique()), float(active["valor_entrada"].sum()),
               float(active["valor_saida"].sum()), float(active["liquido_valor"].sum())],
              style="total", money_cols=(6, 7, 8))

    if summary_only:
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    # 02 — OP (linha visível) -> todas as saídas reais da OP (linhas recolhidas).
    op_headers = ["OP / Código", "Descrição / Situação", "TES", "Movimentos",
                  "Produtos", "Baixas", "Qtd baixada", "UM", "Saída (R$)",
                  "Data / Período", "Armazém", "Documento", "AE original", "Grupo", "Subgrupo"]
    op_sheet = start_sheet("Por OP", op_headers,
                           [26, 68, 12, 16, 15, 12, 17, 12, 23, 28, 14, 25, 26, 32, 41],
                           outline=True)
    op_data = summarize_ops(active)
    if not op_data.empty:
        linked = active.loc[active["op_identificada"] != ""]
        has_saida = (linked["qtd_saida"] != 0) | (linked["valor_saida"] != 0)
        outgoing = linked.loc[has_saida].sort_values(
            ["op_identificada", "data", "codigo", "linha_origem", "id_registro"],
            ascending=[True, False, True, True, True], kind="stable"
        )
        detail_by_op = {key: part for key, part in outgoing.groupby("op_identificada", sort=False)}
        for op in op_data.itertuples(index=False):
            children = detail_by_op.get(op.op_identificada)
            quantity_rows = len(children) if children is not None else 0
            interval = date.fromisoformat(op.primeira_data).strftime("%d/%m/%Y")
            end_date = date.fromisoformat(op.ultima_data).strftime("%d/%m/%Y")
            interval = interval if interval == end_date else interval + " a " + end_date
            write(op_sheet, [
                op.op_identificada, "TOTAL DA OP", "", int(op.lancamentos),
                int(op.produtos), quantity_rows, None, "",
                float(op.valor_saida), interval, "", "", "", "", "",
            ], style="group", collapsed=True, money_cols=(9,))
            if children is None or children.empty:
                write(op_sheet, ["", "SEM BAIXAS NO RECORTE", "", "", "", "",
                                 None, "", 0.0, None, "", "", "", "", ""],
                      style="child", level=1, money_cols=(9,))
                continue
            for movement in children.itertuples(index=False):
                write(op_sheet, [
                    str(movement.codigo), str(movement.descricao), str(movement.tes),
                    None, None, None, float(movement.qtd_saida), str(movement.um),
                    float(movement.valor_saida), movement.data, str(movement.armazem),
                    str(movement.documento), str(movement.ae_original),
                    str(movement.grupo_usuario), str(movement.subgrupo_usuario),
                ], style="child", level=1, date_cols=(10,), money_cols=(9,),
                    quantity_cols=(7,))

    # 03 — Mesmas colunas e ordem do KARDEX > Registros, incluindo subgrupo.
    ledger_headers = ["Data", "Código", "Produto", "Armazém", "TES", "Grupo", "Subgrupo",
                      "Entrada (R$)", "Saída (R$)", "AE original",
                      "Descrição TES", "Tipo do material", "Documento"]
    ledger_sheet = start_sheet(
        "Movimentações", ledger_headers,
        [17, 17, 80, 15, 12, 33, 41, 22, 22, 29, 53, 20, 29]
    )
    if not active.empty:
        ordered = active.sort_values(
            ["data", "codigo", "linha_origem", "id_registro"],
            ascending=[False, True, True, True], kind="stable"
        )
        for row in ordered.itertuples(index=False):
            write(ledger_sheet, [
                row.data, str(row.codigo), str(row.descricao),
                str(row.armazem), str(row.tes),
                str(row.grupo_usuario), str(row.subgrupo_usuario),
                float(row.valor_entrada), float(row.valor_saida),
                str(row.ae_original), str(row.descricao_tes),
                str(row.tipo_analise), str(row.documento),
            ], date_cols=(1,), money_cols=(8, 9))
    ledger_sheet.auto_filter.ref = f"A1:M{max(counters['Movimentações'], 1)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_history(model: dict) -> bytes:
    rows = model["raw_rows"]
    coverage = model.get("coverage_metas") or []
    starts = [m["periodo_inicio"] for m in coverage] or [r["data"] for r in rows]
    ends = [m["periodo_fim"] for m in coverage] or [r["data"] for r in rows]
    manifest = {"formato": HISTORY_FORMAT, "versao_regras": model["versao"], "lancamentos": len(rows),
                "periodo_inicio": min(starts), "periodo_fim": max(ends), "coberturas": coverage,
                "fontes": model["imports"]}
    output = io.BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("manifesto.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        archive.writestr("movimentacoes.jsonl", "".join(json.dumps(r, ensure_ascii=False, separators=(",", ":"))+"\n" for r in rows))
        archive.writestr("regras_classificacao.json", RULES_FILE.read_bytes())
    return output.getvalue()
