"""Relatórios gerenciais do fechamento SETTA.

Exporta uma única planilha por relatório escolhido, sem alterar bases,
regras de conferência ou lançamentos originais.
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import math

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER = "17243B"
LINE = "F2F6FB"
WARNING = "FDE8E8"
MONEY_FORMAT = '#,##0.00;[Red](#,##0.00)'
# "General" respeita a configuração regional do Excel e evita o
# separador decimal isolado que o formato #,##0.###### gerava no pt-BR.
QUANTITY_FORMAT = "General"
QUANTITY_INTEGER_FORMAT = "General"
EXPORT_SCHEMA_VERSION = "2026-10-30.3"
PERCENT_FORMAT = '0.00%;[Red](0.00%)'

OP_FIELDS = [
    ("OP", "op"),
    ("Produtos totais", "itens"),
    ("Produtos separados", "separados"),
    ("Produtos conferidos", "conferidos"),
    ("Produtos baixados", "baixados"),
    ("Produtos com alerta", "divergencias"),
    ("Cliente", "cliente"),
    ("Condição", "condicao"),
    ("Nota fiscal", "nota_fiscal"),
    ("Faturamento", "data_faturamento"),
    ("Pátio", "dias_patio"),
    ("Status da OP", "status_op"),
]
ITEM_FIELDS = [
    ("Produto", "codigo"),
    ("Descrição", "descricao"),
    ("Lotes", "lotes"),
    ("Qtd necessária", "necessaria"),
    ("Qtd separada", "separada"),
    ("Qtd conferida", "conferida"),
    ("Qtd baixada", "baixada"),
    ("Qtd pendente baixa", "pendente"),
    ("Valor previsto da baixa (R$)", "valor_previsto_baixa"),
    ("Saldo no lote conferido", "saldo"),
    ("Saldo G9-PROD", "saldo_g9_prod"),
    ("Outros endereços", "outros_enderecos"),
    ("Resultado", "status"),
]
OP_COLUMNS = [label for label, _ in OP_FIELDS] + [label for label, _ in ITEM_FIELDS]


def _cell(value):
    if value is None:
        return None
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.isoformat()[:10]
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(x) for x in value)
    if isinstance(value, dict):
        return str(value)
    if isinstance(value, str):
        if value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _date(value) -> date | None:
    if value is None or str(value).strip() in ("", "NI", "None"):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value).strip()
    try:
        return date.fromisoformat(raw[:10])
    except ValueError:
        for fmt in ("%d/%m/%Y", "%d/%m/%y"):
            try:
                return datetime.strptime(raw[:10], fmt).date()
            except ValueError:
                pass
    return None


def _period(start: date | None, end: date | None):
    if start is not None and end is not None and start > end:
        raise ValueError("A data inicial não pode ser maior que a final.")


def _between(value, start: date | None, end: date | None) -> bool:
    parsed = _date(value)
    return parsed is not None and (start is None or parsed >= start) and (end is None or parsed <= end)


def ops_with_activity(model: dict, start: date | None, end: date | None) -> set[str]:
    """Seleciona OPs com separação, conferência OU baixa no período.

    Nunca usa data de faturamento como critério implícito. A OP selecionada
    será exportada com todos os seus produtos, mesmo sem evento individual.
    """
    _period(start, end)
    if start is None and end is None:
        return {str(op) for op in model["ops"]["op"].tolist()}
    selected = set()
    for item in model.get("general") or []:
        if any(_between(item.get(field), start, end) for field in ("separacao", "conferencia")):
            selected.add(str(item["op"]))
    for movement in model.get("movements") or []:
        if movement.get("natureza") not in ("BAIXA / DEVOLUÇÃO", "CONSUMO DE OP"):
            continue
        if not _between(movement.get("data"), start, end):
            continue
        if (abs(float(movement.get("qtd_saida") or 0)) > 0
                or abs(float(movement.get("qtd_entrada") or 0)) > 0):
            selected.add(str(movement.get("op") or ""))
    return selected


class _Excel:
    """Escrita streaming com cabeçalho fixo, autofiltro e formatação de colunas."""
    def __init__(self):
        self.workbook = Workbook(write_only=True)
        self.sheets = {}

    def sheet(self, name, headers, *, money=(), quantity=(), percent=(), alert_column=None):
        ws = self.workbook.create_sheet(name[:31])
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        for i, header in enumerate(headers, 1):
            width = min(60, max(16, len(header) + 4))
            if header in ("Descrição", "Outros endereços", "Resultado"):
                width = 46
            ws.column_dimensions[get_column_letter(i)].width = width
        self.sheets[name] = (ws, list(headers), set(money), set(quantity),
                             set(percent), alert_column, 1)
        self.add(name, headers, header=True)
        return ws

    def add(self, name, values, *, header=False, alert=False):
        ws, headers, money, quantities, percents, alert_column, count = self.sheets[name]
        row = []
        for i, value in enumerate(values):
            value = _cell(value)
            c = WriteOnlyCell(ws, value)
            if isinstance(value, str):
                c.data_type = "s"
            if header:
                c.fill = PatternFill("solid", fgColor=HEADER)
                c.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
                c.alignment = Alignment(vertical="center", wrap_text=True)
            else:
                c.font = Font(name="Aptos", size=10, color="17243B")
                if alert and (alert_column is None or alert_column == i):
                    c.fill = PatternFill("solid", fgColor=WARNING)
                elif count % 2 == 0:
                    c.fill = PatternFill("solid", fgColor=LINE)
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    if i in money:
                        c.number_format = MONEY_FORMAT
                    elif i in quantities:
                        # Excel reconhece o valor numérico e determina como
                        # exibir o separador decimal conforme o idioma local.
                        c.number_format = QUANTITY_FORMAT
                    elif i in percents:
                        c.number_format = PERCENT_FORMAT
            row.append(c)
        ws.append(row)
        self.sheets[name] = (ws, headers, money, quantities, percents, alert_column, count + 1)
        if header:
            ws.row_dimensions[1].height = 32

    def save(self) -> bytes:
        for name, (ws, headers, *_rest) in self.sheets.items():
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{self.sheets[name][-1] - 1}"
        output = BytesIO()
        self.workbook.save(output)
        return output.getvalue()


STOCK_HEADERS = ["Competência", "Estoque inicial (R$)", "Estoque final (R$)",
                 "Resumo (R$)", "Variação (%)", "Validação", "Fonte"]


def stock_data(summaries: list[dict], imports: list[dict],
               start: date | None = None, end: date | None = None) -> dict:
    """Reusa o valor_total mensal do Dashboard; não soma meses como saldo."""
    _period(start, end)
    imports_map = {str(x.get("competencia") or "")[:10]: x for x in imports}
    values = {
        (str(r.get("competencia") or "")[:10], str(r.get("dimensao") or ""), str(r.get("chave") or "")):
        float(r.get("valor_total") or 0)
        for r in summaries
    }
    months = sorted(d for key in imports_map if (d := _date(key)) is not None)
    months = [d for d in months if (start is None or d.year * 12 + d.month >= start.year * 12 + start.month)
              and (end is None or d.year * 12 + d.month <= end.year * 12 + end.month)]
    all_month_keys = set(imports_map)
    entries = []
    for month in months:
        key = month.isoformat()
        prev = date(month.year - (month.month == 1), month.month - 1 or 12, 1)
        prior_exists = prev.isoformat() in all_month_keys
        prior_key = prev.isoformat()
        final = values.get((key, "TOTAL", "TOTAL"))
        initial = values.get((prior_key, "TOTAL", "TOTAL")) if prior_exists else None
        if final is None:
            continue
        info = imports_map[key]
        summary = final - initial if initial is not None else None
        change = summary / abs(initial) if initial not in (None, 0) else None
        common = [month.strftime("%m/%Y"), initial, final, summary, change,
                  str(info.get("status") or "NI"), "PRÉVIA" if info.get("preview") else "FECHAMENTO"]
        entries.append((month, common))
    def dimension_rows(dim, prefix):
        results = []
        for month, common in entries:
            key = month.isoformat()
            prev = date(month.year - (month.month == 1), month.month - 1 or 12, 1).isoformat()
            codes = {
                k[2][len(prefix):] for k in values
                if k[1] == dim and k[0] in (key, prev) and k[2].startswith(prefix)
            }
            for code in sorted(codes):
                final = values.get((key, dim, prefix + code), 0.0)
                initial = values.get((prev, dim, prefix + code), 0.0) if prev in all_month_keys else None
                change = final - initial if initial is not None else None
                percent = change / abs(initial) if initial not in (None, 0) else None
                results.append([common[0], code, initial, final, change, percent, common[5], common[6]])
        return results
    return {
        "Resumo Geral": [entries[-1][1]] if entries else [],
        "Por Armazém": dimension_rows("ARMZ", "ARMZ:"),
        "Por Tipo": dimension_rows("TP", "TP:"),
        "Evolução Mensal": [entry for _, entry in entries],
    }


def export_stock(summaries: list[dict], imports: list[dict],
                 start: date | None = None, end: date | None = None) -> bytes:
    data = stock_data(summaries, imports, start, end)
    wb = _Excel()
    for name, rows in data.items():
        headers = STOCK_HEADERS if name in ("Resumo Geral", "Evolução Mensal") else [
            "Competência", "Armazém" if name == "Por Armazém" else "Tipo de material",
            *STOCK_HEADERS[1:],
        ]
        offset = 1 if len(headers) == 8 else 0
        wb.sheet(name, headers, money={i + offset for i in (1, 2, 3)},
                 percent={4 + offset})
        for row in rows:
            wb.add(name, row)
    return wb.save()


def _valor_previsto_baixa(item: dict, prices: dict) -> float | None:
    """Valor estimado somente para material com alerta e pendência positiva.

    Base: preço unitário ULT. PRECO do CADASTROS x quantidade pendente.
    Ausência de preço, divergência de dados ou histórico incompleto não
    devem virar valores monetários artificiais.
    """
    if not bool(item.get("divergente")):
        return None
    try:
        pendente = item.get("pendente")
        codigo = str(item.get("codigo") or "").strip()
        if pendente is None or not codigo:
            return None
        if codigo.isdecimal() and len(codigo) <= 8:
            codigo = codigo.zfill(8)
        preco_bruto = prices.get(codigo)
        if preco_bruto is None or str(preco_bruto).strip() == "":
            return None
        quantidade = Decimal(str(pendente))
        preco_unitario = Decimal(str(preco_bruto))
        if not quantidade.is_finite() or not preco_unitario.is_finite():
            return None
        if quantidade <= 0 or preco_unitario <= 0:
            return None
        return float((quantidade * preco_unitario).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP))
    except (ValueError, ArithmeticError, TypeError):
        return None


def issue_rows(model: dict, start: date | None = None, end: date | None = None,
               *, statuses: set[str] | None = None, alerts_only: bool = False,
               material_types: set[str] | None = None) -> list[list]:
    selected = ops_with_activity(model, start, end)
    op_lookup = {
        str(entry["op"]): entry
        for entry in model["ops"].to_dict("records")
    }
    rows = []
    prices = model.get("catalog_prices") or {}
    for item in model["items"].to_dict("records"):
        op = str(item["op"])
        if op not in selected or op not in op_lookup:
            continue
        head = op_lookup[op]
        if statuses and str(head.get("status_op") or "") not in statuses:
            continue
        if alerts_only and not bool(head.get("divergencias")):
            continue
        if material_types and str(item.get("tipo_material") or "").upper() not in material_types:
            continue
        values = [
            _valor_previsto_baixa(item, prices)
            if field == "valor_previsto_baixa" else item.get(field)
            for _, field in ITEM_FIELDS
        ]
        rows.append([head.get(field) for _, field in OP_FIELDS] + values)
    return rows


def export_issues(model: dict, start: date | None = None, end: date | None = None,
                  *, statuses: set[str] | None = None, alerts_only: bool = False,
                  material_types: set[str] | None = None) -> bytes:
    """Uma única aba, uma linha por produto/OP, dados da OP repetidos por linha."""
    rows = issue_rows(model, start, end, statuses=statuses, alerts_only=alerts_only,
                      material_types=material_types)
    wb = _Excel()
    wb.sheet("Conferência de Baixas", OP_COLUMNS,
             quantity={OP_COLUMNS.index(label) for label in (
                 "Qtd necessária", "Qtd separada", "Qtd conferida",
                 "Qtd baixada", "Qtd pendente baixa",
                 "Saldo no lote conferido", "Saldo G9-PROD")},
             money={OP_COLUMNS.index("Valor previsto da baixa (R$)")},
             alert_column=len(OP_COLUMNS) - 1)
    for row in rows:
        wb.add("Conferência de Baixas", row,
               alert=("NÃO CONSUMIDO" in str(row[-1]).upper()
                      or "BAIXA ACIMA" in str(row[-1]).upper()))
    return wb.save()


def filtered_kardex(frame: pd.DataFrame, start: date | None = None, end: date | None = None,
                    warehouses: set[str] | None = None, types: set[str] | None = None,
                    groups: set[str] | None = None, group_names: set[str] | None = None,
                    group_config: dict | None = None) -> pd.DataFrame:
    _period(start, end)
    filtered = frame
    if start or end:
        dates = pd.to_datetime(frame["data"], errors="coerce").dt.date
        mask = pd.Series(True, index=frame.index)
        if start:
            mask &= dates >= start
        if end:
            mask &= dates <= end
        filtered = filtered.loc[mask]
    if warehouses:
        filtered = filtered.loc[filtered["armazem"].astype(str).isin(warehouses)]
    if types:
        filtered = filtered.loc[filtered["tipo"].astype(str).str.upper().isin(types)]
    if groups:
        filtered = filtered.loc[filtered["tes"].astype(str).str.zfill(3).isin(groups)]
    if group_names:
        if group_config is None:
            raise ValueError("O filtro de grupos exige o cadastro gerencial atualizado.")
        import kardex_group_config as gc
        owner = gc.group_owner(group_config)
        named = filtered["tes"].astype(str).str.zfill(3).map(owner).fillna(gc.UNASSIGNED)
        filtered = filtered.loc[named.isin(group_names)]
    return filtered.copy()


def export_kardex(frame: pd.DataFrame, group_config: dict, *,
                  summary_only: bool = False) -> bytes:
    import movements
    return movements.export_excel(frame, group_config=group_config, summary_only=summary_only)
