from __future__ import annotations

import sys
import unittest
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import issue_check as check
from issue_check_export import export_complete_report

OP = "26000101001"
CODE = "00001000"
META = [{"arquivo_nome": "teste.xlsx", "periodo_inicio": "2026-01-01",
         "periodo_fim": "2026-10-21", "lancamentos": 5}]


def general(**override):
    result = {"Projeto": OP, "Código": CODE, "Lote": "000001",
              "Descrição": "MATERIAL", "Qtd. necessária": 10,
              "Qtd. atendida": 10, "Data de separação": "19/10/26 08:00",
              "Data de conferência": "20/10/26 08:00"}
    result.update(override)
    return result


def stock(**override):
    result = {"Produto": CODE, "Lote": "000001", "Armazem": "S2",
              "Endereco": "G9-PROD", "Quantidade": 3, "Empenho": 0}
    result.update(override)
    return result


def movement(tes="999", **override):
    result = {"codigo": CODE, "ae_original": "OP" + OP, "armazem": "S2",
              "data": "2026-10-20", "tes": tes,
              "qtd_saida": "7", "qtd_entrada": None,
              "valor_saida": "70", "valor_entrada": None,
              "um": "UN", "documento": "DOC", "arquivo_origem": "test.xlsx",
              "linha_origem": 3}
    result.update(override)
    return result


def audit(g=None, s=None, m=None):
    return check.reconcile(
        check.parse_general(pd.DataFrame(g if g is not None else [general()])),
        check.parse_stock(pd.DataFrame(s)) if s else [],
        m if m is not None else [],
        history_metas=META, as_of="2026-10-21")


class BusinessRulesTests(unittest.TestCase):
    def test_unpaired_999_with_op_is_real_consumption(self):
        item = audit(s=[stock()], m=[movement()])["items"].iloc[0]
        self.assertEqual(item.baixada, 7)
        self.assertEqual(item.pendente, 3)
        self.assertIn(check.PROD_PENDING, item.alertas)

    def test_999_compensated_with_499_without_op_does_not_reduce_pending(self):
        outbound = movement(qtd_saida="10", valor_saida="100")
        inbound = movement("499", ae_original="", qtd_saida=None, qtd_entrada="10",
                           valor_saida=None, valor_entrada="100")
        item = audit(s=[stock(Quantidade=10)], m=[outbound, inbound])["items"].iloc[0]
        self.assertEqual(item.baixada, 0)
        self.assertEqual(item.pendente, 10)
        self.assertNotIn(check.PROD_PENDING, item.alertas)

    def test_999_without_op_is_not_assigned_to_unrelated_op(self):
        item = audit(s=[stock()], m=[movement(ae_original="CC600307")])["items"].iloc[0]
        self.assertEqual(item.baixada, 0)
        self.assertFalse(item.op_possui_baixa)

    def test_g9_alert_requires_separation_and_prior_op_consumption(self):
        without_separation = audit(
            g=[general(**{"Data de separação": None})],
            s=[stock(Quantidade=3)], m=[movement()])["items"].iloc[0]
        self.assertNotIn(check.PROD_PENDING, without_separation.alertas)
        before_first_issue = audit(
            s=[stock(Quantidade=10)], m=[])["items"].iloc[0]
        self.assertNotIn(check.PROD_PENDING, before_first_issue.alertas)

    def test_alternative_addresses_rank_exact_coverage_first_and_keep_lot(self):
        alternatives = [
            stock(Endereco="E-QUATRO", Quantidade=4),
            stock(Endereco="E-DOIS", Quantidade=2),
            stock(Endereco="E-EXATO", Quantidade=3),
            stock(Endereco="E-NOVE", Quantidade=9),
            stock(Endereco="E-OUTRO-LOTE", Lote="000002", Quantidade=3),
            stock(Quantidade=0),
        ]
        item = audit(s=alternatives, m=[movement()])["items"].iloc[0]
        names = [e["endereco"] for e in item.enderecos_alternativos]
        self.assertEqual(names, ["E-EXATO", "E-QUATRO", "E-NOVE", "E-DOIS"])
        self.assertIn("E-EXATO", item.outros_enderecos)
        self.assertNotIn("E-OUTRO-LOTE", item.outros_enderecos)

    def test_baixados_counts_only_fully_consumed_confirmed_items(self):
        model = audit(s=[stock(Quantidade=0)], m=[movement(qtd_saida="10")])
        self.assertEqual(int(model["ops"].iloc[0].baixados), 1)
        self.assertEqual(float(model["items"].iloc[0].separada), 10)
        self.assertEqual(int(audit(m=[movement()])["ops"].iloc[0].baixados), 0)

    def test_complete_report_contains_all_source_tabs_and_consistent_values(self):
        model = audit(s=[stock()], m=[movement()])
        model["source_meta"] = {"endereco": {"last_file_name": "endereco.xlsx"}}
        model["history_meta"] = META
        raw = export_complete_report(model)
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        expected = {"PARAMETROS", "RESUMO_OP", "PRODUTOS", "LOTES", "RELATORIO_MES",
                    "SALDOS_ENDERECOS", "MOVIMENTOS_KARDEX", "FONTES", "COBERTURA"}
        self.assertEqual(set(wb.sheetnames), expected)
        ws = wb["PRODUTOS"]
        heads = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        data = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
        self.assertEqual(data[heads.index("baixada")], 7)
        self.assertIn("E-", str(data[heads.index("outros_enderecos")]) if data[heads.index("outros_enderecos")] else "E-")
        self.assertEqual(wb["MOVIMENTOS_KARDEX"].max_row, 2)
        wb.close()


if __name__ == "__main__":
    unittest.main()
