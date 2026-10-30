"""Regressão dos três relatórios e filtro de atividade da OP."""
from __future__ import annotations

from datetime import date
from io import BytesIO
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import reports


def op(n, **overrides):
    base = {
        "op": n, "itens": 2, "separados": 2, "conferidos": 1,
        "baixados": 0, "divergencias": 1, "cliente": "CLIENTE X",
        "condicao": "NI", "nota_fiscal": "000019", "data_faturamento": "NI",
        "dias_patio": "NI", "status_op": "EM CONFERÊNCIA",
    }
    base.update(overrides)
    return base


def item(n, code, kind="MP", result="PRODUTO NÃO CONSUMIDO"):
    return {
        "op": n, "codigo": code, "descricao": "MATERIAL", "lotes": "00000000",
        "necessaria": 11, "separada": 11, "conferida": 11,
        "baixada": 0, "pendente": 11, "saldo": 10,
        "saldo_g9_prod": 8, "outros_enderecos": "",
        "status": result, "tipo_material": kind,
    }


class ReportsTests(unittest.TestCase):
    def setUp(self):
        self.model = {
            "ops": pd.DataFrame([
                op("0001"), op("0002"), op("0003"), op("0004"),
            ]),
            "items": pd.DataFrame([
                item("0001", "00006001"), item("0001", "00006002", kind="II", result="MATERIAL DE CONSUMO"),
                item("0002", "00003002"), item("0003", "00003003"),
                item("0004", "00004004"),
            ]),
            "general": [
                {"op":"0001", "separacao":"2026-09-25T08:00:00", "conferencia":"2026-10-03T08:00:00"},
                {"op":"0001", "separacao":"2026-09-25T08:00:00", "conferencia":""},
                {"op":"0002", "separacao":"2026-10-02T09:00:00", "conferencia":""},
                {"op":"0003", "separacao":"2026-09-01T09:00:00", "conferencia":"2026-09-02T09:00:00"},
                {"op":"0004", "separacao":"2026-09-20T09:00:00", "conferencia":""},
            ],
            "movements": [
                {"op":"0003", "data":"2026-10-07", "natureza":"CONSUMO DE OP",
                 "qtd_saida":2, "qtd_entrada":0},
                {"op":"0004", "data":"2026-10-07", "natureza":"ENDEREÇAMENTO",
                 "qtd_saida":2, "qtd_entrada":0},
            ],
        }

    def test_activity_includes_whole_op_but_not_unrelated_op(self):
        start, end = date(2026, 10, 1), date(2026, 10, 10)
        self.assertEqual(reports.ops_with_activity(self.model,start,end),
                         {"0001","0002","0003"})
        rows=reports.issue_rows(self.model,start,end)
        self.assertEqual(len(rows), 4)
        self.assertEqual([r[0] for r in rows], ["0001","0001","0002","0003"])
        self.assertEqual(rows[1][len(reports.OP_FIELDS)], "00006002")
        self.assertEqual(len(reports.issue_rows(self.model)), 5)

    def test_issue_excel_is_one_sheet_columns_flat_and_codes_text(self):
        raw=reports.export_issues(self.model,date(2026,10,1),date(2026,10,10))
        wb=load_workbook(BytesIO(raw),read_only=True,data_only=True)
        try:
            self.assertEqual(wb.sheetnames,["Conferência de Baixas"])
            values=list(wb.active.values)
            self.assertEqual(len(values),5)
            self.assertEqual(values[0],tuple(reports.OP_COLUMNS))
            self.assertEqual(values[1][0],"0001")
            self.assertEqual(values[1][12],"00006001")
            self.assertEqual(values[2][12],"00006002")
            self.assertEqual(values[1][8],"000019")
        finally:
            wb.close()

    def test_quantity_formats_do_not_append_decimal_separator_to_integers(self):
        wb = reports._Excel()
        wb.sheet("Quantidades", ["Valor"], quantity={0})
        for value in (100, 100.0, 0, -2, 3.25):
            wb.add("Quantidades", [value])
        raw = wb.save()
        saved = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        try:
            sheet = saved["Quantidades"]
            self.assertEqual([sheet[f"A{n}"].value for n in range(2, 7)],
                             [100, 100, 0, -2, 3.25])
            for number in range(2, 6):
                self.assertEqual(sheet[f"A{number}"].number_format,
                                 reports.QUANTITY_INTEGER_FORMAT)
            self.assertEqual(sheet["A6"].number_format, reports.QUANTITY_FORMAT)
        finally:
            saved.close()

    def test_issue_report_quantity_is_formatted_as_integer(self):
        raw = reports.export_issues(self.model)
        saved = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        try:
            # Coluna P: quantidade necessária. A célula deve exibir 11, não 11,.
            self.assertEqual(saved.active["P2"].value, 11)
            self.assertEqual(saved.active["P2"].number_format,
                             reports.QUANTITY_INTEGER_FORMAT)
            # Código do material continua sendo texto (zeros à esquerda).
            self.assertEqual(saved.active["M2"].value, "00006001")
        finally:
            saved.close()

    def test_type_and_alert_filters_preserve_op_activity(self):
        rows=reports.issue_rows(self.model,date(2026,10,1),date(2026,10,10),
                                material_types={"II"},alerts_only=True)
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0][12],"00006002")

    def test_stock_workbook_has_four_tabs_and_no_fake_initial_balance(self):
        entries=[]
        imports=[]
        for month, total in ((8,100.0),(9,110.0),(10,120.0)):
            day=date(2026,month,1).isoformat()
            imports.append({"competencia":day,"status":"VALIDO","preview":month==10})
            entries.extend([
                {"competencia":day,"dimensao":"TOTAL","chave":"TOTAL","valor_total":total},
                {"competencia":day,"dimensao":"ARMZ","chave":"ARMZ:S2","valor_total":total},
                {"competencia":day,"dimensao":"TP","chave":"TP:MP","valor_total":total},
            ])
        result=reports.stock_data(entries,imports,date(2026,9,15),date(2026,10,20))
        self.assertEqual(len(result["Evolução Mensal"]),2)
        self.assertEqual(result["Evolução Mensal"][0][1:4],[100.0,110.0,10.0])
        self.assertEqual(result["Evolução Mensal"][1][-1],"PRÉVIA")
        self.assertEqual(result["Por Armazém"][0][1],"S2")
        raw=reports.export_stock(entries,imports)
        wb=load_workbook(BytesIO(raw),read_only=True,data_only=True)
        try:
            self.assertEqual(wb.sheetnames,["Resumo Geral","Por Armazém","Por Tipo","Evolução Mensal"])
            self.assertIsNone(list(wb["Evolução Mensal"].values)[1][1])
        finally:
            wb.close()

    def test_kardex_filters_are_non_destructive(self):
        frame=pd.DataFrame([
            {"data":"2026-10-03","armazem":"S2","tipo":"II","tes":"520"},
            {"data":"2026-10-05","armazem":"EP","tipo":"EP","tes":"520"},
            {"data":"2026-09-30","armazem":"S2","tipo":"MP","tes":"020"},
        ])
        result=reports.filtered_kardex(frame,date(2026,10,1),date(2026,10,31),
                                       warehouses={"S2"},types={"II"},groups={"520"})
        self.assertEqual(len(result),1)
        self.assertEqual(len(frame),3)
        group_result=reports.filtered_kardex(
            frame,group_names={"FLUXO DE ESTOQUE"},
            group_config={"grupos":[{"nome":"FLUXO DE ESTOQUE","tes":["520"]}]})
        self.assertEqual(len(group_result),2)
        with patch("movements.export_excel",return_value=b"excel") as fake:
            self.assertEqual(reports.export_kardex(result,{"condicoes":[]},summary_only=True),b"excel")
            self.assertTrue(fake.call_args.kwargs["summary_only"])


if __name__=="__main__":
    unittest.main()
