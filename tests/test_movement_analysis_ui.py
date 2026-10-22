from __future__ import annotations

import io
import sys
import unittest
from pathlib import Path

import openpyxl
from streamlit.testing.v1 import AppTest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import movements as mv


PROGRAM = """
import sys
import streamlit as st
sys.path.insert(0, "tests")
from test_movement_analysis import record, imported
import movements as mv
import movement_analysis_ui as ui
import kardex_group_config as group_cfg
ui._groups_read = lambda: group_cfg.empty()

rows = [
    record(documento="AJUSTE"),
    record(tes="550", ae_original=" OP00001234567 ", documento="OP", valor_saida="20"),
    record(tes="020", tipo="PA", documento="PRODUCAO", qtd_saida=None, valor_saida=None,
           qtd_entrada="3", valor_entrada="30"),
    record(tes="001", documento="SETEMBRO", data="2026-09-10", qtd_saida=None, valor_saida=None,
           qtd_entrada="4", valor_entrada="40"),
    record(tes="550", ae_original="OP00007654321", documento="OUTRA OP", armazem="S1", valor_saida="5"),
]
pack = imported(rows)
pack["meta"].update(periodo_inicio="2026-09-01", periodo_fim="2026-10-19")
model = mv.build_model([pack])
imports = [{"chave": "synthetic", "meta": {"arquivo_nome": "teste.xlsx", "periodo_inicio": "2026-09-01",
            "periodo_fim": "2026-10-19", "lancamentos": 5}, "atualizado_em": "2026-10-21T10:00:00Z"}]
action = st.session_state.pop("_test_action", None)
if action == "tes":
    st.session_state.test_selection = {"selection": {"rows": [0]}}
    ui._tes_selected("test_selection", ["520"])
elif action == "op":
    st.session_state.test_selection = {"selection": {"rows": [0]}}
    ui._op_selected("test_selection", ["OP00001234567"])
elif action in ("group_only", "drill"):
    original_table = ui._table
    def select_in_order(data, *args, **kwargs):
        displayed = original_table(data, *args, **kwargs)
        key = kwargs.get("key", "")
        if "_overview_groups_" in key:
            return {"selection": {"rows": [0]}}
        if action == "drill" and ("_overview_subgroups_" in key or "_overview_tes_" in key):
            return {"selection": {"rows": [0]}}
        return displayed
    ui._table = select_in_order
try:
    ui.render_results(model, imports)
finally:
    if action in ("group_only", "drill"):
        # Streamlit AppTest reutiliza o módulo entre execuções; não deixar
        # o mock de seleção interferir nos demais testes.
        ui._table = original_table
"""


class MovementAnalysisUITests(unittest.TestCase):
    def app(self):
        app = AppTest.from_string(PROGRAM, default_timeout=20).run()
        self.assertFalse(app.exception)
        return app

    def area(self, app, name):
        app.segmented_control[0].set_value(name).run()
        self.assertFalse(app.exception)

    def test_all_visible_kardex_dates_follow_ptbr(self):
        from datetime import date
        import movement_analysis_ui as ui
        self.assertEqual(ui._br_date("2026-10-22"), "22/10/2026")
        self.assertEqual(ui._br_date(date(2026, 1, 5)), "05/01/2026")
        self.assertEqual(ui._br_month("2026-10"), "10/2026")
        app = self.app()
        self.area(app, "Por OP")
        op = app.dataframe[0].value
        self.assertIn("Primeiro movimento", op.columns)
        self.assertTrue(op["Primeiro movimento"].str.match(r"^\d{2}/\d{2}/\d{4}$").all())
        self.assertTrue(op["Último movimento"].str.match(r"^\d{2}/\d{2}/\d{4}$").all())
        self.area(app, "Movimentações")
        ledger = app.dataframe[0].value
        self.assertTrue(ledger["Data"].notna().all())
        self.assertIn("Grupo", ledger.columns)
        self.assertIn("Subgrupo", ledger.columns)

    def test_all_registered_groups_display_even_without_movements(self):
        from test_movement_analysis import imported, record
        import movement_analysis_ui as ui
        import kardex_group_config as cg
        import kardex_conditions as cond
        sample = mv.build_model([imported([record(tes="550")])])["frame"]
        cfg = cg.create_group(cg.empty(), "COMPRAS", known_tes={"550","020"})
        cfg = cg.create_group(cfg, "FLUXO DE ESTOQUE", known_tes={"550","020"})
        cfg = cg.assign_group(cfg, cfg["grupos"][1]["id"], ["550"],
                              known_tes={"550","020"})
        classified = cond.classify(sample, cfg)
        result = ui._overview_group_totals(classified, cfg)
        self.assertEqual(result["Grupo"].tolist(), ["COMPRAS", "FLUXO DE ESTOQUE"])
        self.assertEqual(result["Movimentos"].tolist() if "Movimentos" in result
                         else result["lancamentos"].tolist(), [0, 1])
        self.assertEqual(result["valor_entrada"].sum(), sample["valor_entrada"].sum())
        self.assertEqual(result["valor_saida"].sum(), sample["valor_saida"].sum())

    def test_dynamic_table_height_avoids_phantom_rows(self):
        import movement_analysis_ui as ui
        self.assertEqual(ui._grid_height(1), 74)
        self.assertEqual(ui._grid_height(2), 109)
        self.assertEqual(ui._grid_height(6), 249)
        self.assertEqual(ui._grid_height(0), 72)
        self.assertEqual(ui._grid_height(30, 550), 550)
        self.assertIsNone(ui._picked_index({"selection": {"rows": []}}, 6))
        self.assertIsNone(ui._picked_index({"selection": {"rows": [6]}}, 6))
        self.assertEqual(ui._picked_index({"selection": {"rows": [2]}}, 6), 2)

    def test_summary_loads_subgroups_only_after_group_checkbox(self):
        app = self.app()
        self.assertFalse(any("Subgrupo" in frame.value.columns for frame in app.dataframe))
        app.session_state["_test_action"] = "group_only"
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(app.segmented_control[0].value, "Visão geral")
        self.assertEqual(app.dataframe[0].value["Grupo"].tolist(), ["SEM GRUPO"])
        self.assertIn("Subgrupo", app.dataframe[1].value.columns)
        self.assertFalse(any("TES" in frame.value.columns for frame in app.dataframe[2:]))

    def test_summary_drills_group_subgroup_tes_movements_without_navigation(self):
        app = self.app()
        app.session_state["_test_action"] = "drill"
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(app.segmented_control[0].value, "Visão geral")
        self.assertGreaterEqual(len(app.dataframe), 4)
        groups = app.dataframe[0].value
        subgroups = app.dataframe[1].value
        tes = app.dataframe[2].value
        ledger = app.dataframe[3].value
        self.assertEqual(groups["Grupo"].tolist(), ["SEM GRUPO"])
        self.assertIn("Subgrupo", subgroups.columns)
        self.assertIn("TES", tes.columns)
        self.assertIn("Grupo", ledger.columns)
        self.assertIn("Subgrupo", ledger.columns)
        self.assertEqual(set(ledger["TES"]), {tes.iloc[0]["TES"]})
        self.assertEqual(set(ledger["Subgrupo"]), {subgroups.iloc[0]["Subgrupo"]})
        self.assertTrue(ledger["Data"].notna().all())

    def test_warehouse_registry_remains_accessible(self):
        from test_movement_analysis import imported, record
        import movement_analysis_ui as ui
        model = mv.build_model([imported([
            record(armazem="S2", documento="S2"),
            record(armazem="E2", data="2026-10-20", documento="E2"),
        ])])
        states = ui._warehouse_registry_overview(model).set_index("Armazém")
        self.assertEqual(states.loc["S2", "Situação"], "ATIVO")
        self.assertEqual(states.loc["E2", "Situação"], "ATIVO")
        self.assertEqual(states.loc["EP", "Situação"], "PENDENTE DE IMPORTAÇÃO")
        self.assertEqual(states.loc["20", "Situação"], "PENDENTE DE IMPORTAÇÃO")

    def test_summary_uses_tes_and_literal_description_only(self):
        app = self.app()
        self.assertEqual([x.label for x in app.multiselect], ["Armazéns"])
        self.assertEqual(app.segmented_control[0].value, "Visão geral")
        frame = app.dataframe[0].value
        self.assertIn("Grupo", frame.columns)
        self.assertEqual(frame["Grupo"].tolist(), ["SEM GRUPO"])
        # Escopo inicial é outubro: o lançamento de setembro não entra nos cards.
        self.assertEqual(frame["Entrada (R$)"].sum(), 30)
        self.assertEqual(frame["Saída (R$)"].sum(), 35)
        self.assertEqual(frame["Movimentos"].sum(), 4)

    def test_tes_local_filter_survives_navigation(self):
        app = self.app()
        self.area(app, "TES e descrições")
        tes = next(x for x in app.multiselect if x.label == "Códigos TES")
        tes.set_value(["520"]).run()
        self.assertFalse(app.exception)
        self.assertEqual(app.dataframe[0].value["TES"].tolist(), ["520"])
        self.area(app, "Por OP")
        next(x for x in app.text_input if x.label == "Buscar OP ou material").set_value("00001234567").run()
        self.assertEqual(app.dataframe[0].value["OP"].tolist(), ["OP00001234567"])
        self.area(app, "Visão geral")
        self.assertEqual(app.dataframe[0].value["Saída (R$)"].sum(), 35)
        self.area(app, "TES e descrições")
        self.assertEqual(next(x for x in app.multiselect if x.label == "Códigos TES").value, ["520"])

    def test_tes_selection_opens_only_matching_ledger_records(self):
        app = self.app()
        app.session_state["_test_action"] = "tes"
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(app.segmented_control[0].value, "Movimentações")
        self.assertEqual(app.dataframe[0].value["TES"].tolist(), ["520"])
        self.assertEqual(len(app.dataframe[0].value), 1)
        self.assertTrue(any("TES: 520" in x.value for x in app.caption))
        next(x for x in app.button if "VOLTAR PARA" in x.label).click().run()
        self.assertFalse(app.exception)
        self.assertEqual(app.segmented_control[0].value, "TES e descrições")

    def test_op_drilldown_preserves_original_ae_and_clears_search(self):
        app = self.app()
        self.area(app, "Movimentações")
        self.assertNotIn("Endereçamento compensado", app.dataframe[0].value.columns)
        next(x for x in app.text_input if x.label == "Buscar nos registros").set_value("INEXISTENTE").run()
        self.assertEqual(len(app.dataframe), 0)
        app.session_state["_test_action"] = "op"
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.dataframe[0].value), 1)
        self.assertEqual(app.dataframe[0].value["AE original"].tolist(), [" OP00001234567 "])
        self.assertEqual(next(x for x in app.text_input if x.label == "Buscar nos registros").value, "")
        next(x for x in app.button if x.label == "LIMPAR CONSULTA").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.dataframe[0].value), 4)

    def test_catalog_independent_of_search_and_sources_accessible(self):
        app = self.app()
        self.area(app, "Movimentações")
        app.text_input[0].set_value("INEXISTENTE").run()
        self.assertEqual(len(app.dataframe), 0)
        self.area(app, "Cadastros")
        # Catálogo: 342 códigos cadastrados mais eventuais TES do KARDEX
        # sem correspondência; nunca descartar os códigos não cadastrados.
        catalog_view = app.dataframe[0].value
        self.assertGreaterEqual(len(catalog_view), 342)
        self.assertEqual(len(catalog_view["TES"].unique()), len(catalog_view))
        self.assertIn("Descrição TES", catalog_view.columns)
        self.assertNotIn("Tipo de TM", catalog_view.columns)
        self.assertNotIn("Origem", catalog_view.columns)
        self.assertEqual(
            catalog_view.loc[catalog_view["TES"] == "540", "Descrição TES"].iloc[0],
            "BAIX CUSTO E QTD ZER",
        )
        self.assertEqual(
            catalog_view.loc[catalog_view["TES"] == "499", "Descrição TES"].iloc[0],
            "ENTRADA DE INVENTÁRIO/ENTRADA OP",
        )
        self.assertEqual(
            catalog_view.loc[catalog_view["TES"] == "999", "Descrição TES"].iloc[0],
            "SAÍDA DE INVENTÁRIO/CONSUMO OP",
        )
        self.assertTrue(set(mv.tes_reference()).issubset(set(catalog_view["TES"])))
        self.assertNotIn("Condição", app.dataframe[0].value.columns)
        self.area(app, "Fontes e relatórios")
        self.assertEqual(len(app.file_uploader), 1)
        self.assertTrue(any("Situação" in frame.value.columns for frame in app.dataframe))

    def test_period_and_warehouse_scope(self):
        app = self.app()
        next(x for x in app.multiselect if x.label == "Armazéns").set_value(["S2"]).run()
        self.assertEqual(app.dataframe[0].value["Saída (R$)"].sum(), 30)
        self.area(app, "Por OP")
        self.assertEqual(app.dataframe[0].value["OP"].tolist(), ["OP00001234567"])
        self.area(app, "Visão geral")
        next(x for x in app.selectbox if x.label == "Período da análise").set_value("ACUMULADO DO ANO").run()
        self.assertEqual(app.dataframe[0].value["Entrada (R$)"].sum(), 70)
        next(x for x in app.selectbox if x.label == "Competência analisada").set_value("2026-09").run()
        self.assertEqual(app.dataframe[0].value["Entrada (R$)"].sum(), 40)

    def test_conditions_priority_visible_in_cadastros(self):
        app = self.app()
        self.area(app, "Cadastros")
        self.assertFalse(app.exception)
        values = [item.value for item in app.markdown]
        self.assertTrue(any("CONDIÇÕES E SUBGRUPOS" in value for value in values))
        self.assertTrue(any("PRODUÇÃO INTERNA" in value for value in values))
        self.assertTrue(any("CONSUMO OP" in value for value in values))
        self.assertTrue(any("ADEQUAÇÃO DE CUSTO" in value for value in values))

    def test_excel_of_filtered_op_preserves_ae(self):
        app = self.app()
        app.session_state["_test_action"] = "op"
        app.run()
        next(x for x in app.button if x.label == "PREPARAR EXCEL DO RECORTE").click().run()
        self.assertFalse(app.exception)
        wb = openpyxl.load_workbook(io.BytesIO(app.session_state["fm_mov_excel"][1]), read_only=True)
        self.assertEqual(wb.sheetnames, ["Grupos e TES", "Por OP", "Movimentações"])
        self.assertEqual(wb["Movimentações"]["F1"].value, "Grupo")
        self.assertEqual(wb["Movimentações"]["G1"].value, "Subgrupo")
        self.assertEqual(wb["Movimentações"]["F2"].value, "SEM GRUPO")
        self.assertEqual(wb["Movimentações"]["G2"].value, "CONSUMO OP")
        self.assertEqual(wb["Movimentações"]["J2"].value, " OP00001234567 ")
        self.assertEqual(len(list(wb["Movimentações"].iter_rows(values_only=True))), 2)
        self.assertEqual(wb["Por OP"]["A2"].value, "OP00001234567")
        self.assertEqual(wb["Grupos e TES"]["A2"].value, "SEM GRUPO")
        wb.close()


if __name__ == "__main__":
    unittest.main()
